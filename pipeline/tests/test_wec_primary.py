"""November candidates come only from certified primary totals and approved
independents; anything the records cannot settle stays open, never guessed."""

import json

import pytest
from openpyxl import Workbook

from importer import import_wec
from importer.checks import check_ballot
from importer.wec_primary import nominee, parse

OFFICE = "REPRESENTATIVE TO THE ASSEMBLY DISTRICT 7"


def sheet(wb, title, names, wards, office_totals=None, cycle=2026):
    """One contest in the Commission's layout: names at fixed columns with a gap."""
    ws = wb.create_sheet()
    ws.append(["WEC Canvass Reporting System"])
    ws.append([f"{cycle} Partisan Primary"])
    ws.append([title])
    cols = [3 + 2 * i for i in range(len(names))]  # a blank column between candidates
    header, name_row = ["", None, "Total Votes Cast"], [None, None, None]
    for c, n in zip(cols, names, strict=True):
        header += [None] * (c - len(header)) + ["DEM"]
        name_row += [None] * (c - len(name_row)) + [n]
    ws.append(header)
    ws.append(name_row)
    totals = [0] * len(names)
    for i, votes in enumerate(wards):
        row = ["MILWAUKEE " if i == 0 else None, f"Ward {i + 1}", sum(map(int, votes))]
        for c, v in zip(cols, votes, strict=True):
            row += [None] * (c - len(row)) + [v]
        ws.append(row)
        totals = [a + int(b) for a, b in zip(totals, votes, strict=True)]
    for label, values in (("County Totals:", totals), ("Office Totals:", office_totals or totals)):
        row = [label, None, sum(totals)]
        for c, v in zip(cols, values, strict=True):
            row += [None] * (c - len(row)) + [v]
        ws.append(row)


def workbook(tmp_path, contests, trailing=None):
    wb = Workbook()
    wb.active.title = "Document map"
    wb.active.append(["Ward by Ward Report"])
    for title, *args in contests:
        wb.active.append([title])
        sheet(wb, title, *args)
    tail = wb.create_sheet()
    tail.append(["WEC Canvass Reporting System"])
    if trailing is not None:
        tail.append(trailing)
    path = tmp_path / "primary.xlsx"
    wb.save(path)
    return path


def test_votes_follow_their_name_columns_and_match_certified_totals(tmp_path):
    path = workbook(tmp_path, [(f"{OFFICE} - Democratic", ["Ann Lee", "Bo Ray", "SCATTERING"],
                                [[10, 3, 1], [5, 9, 0]])])
    (contest,) = parse(path, 2026)
    assert contest["candidates"] == [("Ann Lee", 15), ("Bo Ray", 12)]
    assert (contest["write_in"], contest["cast"]) == (1, 28)
    assert nominee(contest) == ("nominee", "Ann Lee")


@pytest.mark.parametrize("defect", ["totals", "text", "trailing", "cycle"])
def test_any_count_the_file_cannot_prove_stops_the_run(tmp_path, defect):
    names, wards = ["Ann Lee", "SCATTERING"], [[4, 0], [6, 1]]
    if defect == "text":
        wards[1][0] = "6"
    path = workbook(
        tmp_path, [(f"{OFFICE} - Democratic", names, wards,
                    [11, 1] if defect == "totals" else None, 2025 if defect == "cycle" else 2026)],
        trailing=[None, None, 5] if defect == "trailing" else None)
    with pytest.raises(RuntimeError):
        parse(path, 2026)


@pytest.mark.parametrize("candidates,write_in,office,expected", [
    ([("Ann Lee", 50), ("Bo Ray", 50)], 0, OFFICE, "unresolved"),       # tie
    ([("Ann Lee", 50)], 50, OFFICE, "unresolved"),  # unnamed write-ins could lead
    ([("Ann Lee", 40), ("Cy Po (write-in)", 60)], 0, OFFICE, "nominee"),  # 8.16(1): most votes
    ([("Cy Po (Write-in)", 199)], 0, OFFICE, "none"),                 # 8.16(2): below 200
    ([], 199, OFFICE, "none"),
    ([("Cy Po (Write-in)", 200)], 0, OFFICE, "unresolved"),           # may have qualified
    ([], 1999, "GOVERNOR", "none"),
    ([], 2000, "GOVERNOR", "unresolved"),
])
def test_nominees_follow_section_8_16(candidates, write_in, office, expected):
    contest = {"office": office, "party": "Democratic", "candidates": candidates,
               "write_in": write_in, "cast": 0}
    assert nominee(contest)[0] == expected


def contests(**parties):
    return [{"office": OFFICE, "party": party, "candidates": names, "write_in": 0, "cast": 0}
            for party, names in parties.items()]


def report(*rows):
    return [{"office": OFFICE, "incumbent": "", "incumbent_noncandidacy": "0",
             "candidate": name, "party": party, "ballot_status": status}
            for name, party, status in rows]


@pytest.fixture
def rulings(tmp_path, monkeypatch):
    def write(**cycle):
        path = tmp_path / "rulings.json"
        path.write_text(json.dumps({"2026": {"candidates": [], "contests": [], **cycle}}))
        monkeypatch.setattr(import_wec, "RULINGS_PATH", path)
    write()
    return write


def test_november_holds_nominees_and_approved_independents_only(rulings):
    ballot, pending = import_wec.november(
        report(("Ann Lee", "Democratic", "Approve"), ("Bo Ray", "Democratic", "Approve"),
               ("Di Vo", "Republican", "Challenged"), ("Ed Um", "The Common", "Approve"),
               ("Fay Oz", "Olive Party", "Deny")),
        contests(Democratic=[("Ann Lee", 90), ("Bo Ray", 80)],
                 Republican=[("Cy Po (Write-in)", 250)]), 2026)
    assert ballot[OFFICE] == [{"candidate": "Ann Lee", "party": "Democratic"},
                              {"candidate": "Ed Um", "party": "Independent"}]
    assert [party for party, _ in pending[OFFICE]] == ["Republican"]


def test_a_ruling_settles_what_the_records_cannot(rulings):
    rows = report(("Ed Um", "Independent", "Challenged"))
    races = contests(Democratic=[("Cy Po (Write-in)", 250)])
    with pytest.raises(RuntimeError, match="challenge outcome"):
        import_wec.november(rows, races, 2026)
    basis = "https://elections.wi.gov/minutes.pdf"
    rulings(candidates=[{"office": OFFICE, "candidate": "Ed Um", "on_ballot": True,
                         "basis": basis}],
            contests=[{"office": OFFICE, "party": "Democratic", "nominee": "Cy Po (Write-in)",
                       "basis": basis}])
    ballot, pending = import_wec.november(rows, races, 2026)
    assert [e["candidate"] for e in ballot[OFFICE]] == ["Cy Po (Write-in)", "Ed Um"]
    assert not pending


@pytest.mark.parametrize("ruling,match", [
    ({"candidates": [{"office": OFFICE, "candidate": "Ghost", "on_ballot": True,
                      "basis": "https://elections.wi.gov/x.pdf"}]}, "no longer needs"),
    ({"contests": [{"office": OFFICE, "party": "Democratic", "nominee": None,
                    "basis": "https://elections.wi.gov/x.pdf"}]}, "settled primary"),
    ({"contests": [{"office": OFFICE, "party": "Democratic", "nominee": None,
                    "basis": "a news story"}]}, "official basis"),
])
def test_rulings_the_data_does_not_need_stop_the_run(rulings, ruling, match):
    rulings(**ruling)
    with pytest.raises(RuntimeError, match=match):
        import_wec.november(report(("Ann Lee", "Democratic", "Approve")),
                            contests(Democratic=[("Ann Lee", 90)]), 2026)


def test_a_primary_missing_an_office_or_party_stops_the_run(rulings):
    with pytest.raises(RuntimeError, match="does not cover"):
        import_wec.november(report(("Ann Lee", "Democratic", "Approve")),
                            contests(Democratic=[("Ann Lee", 90)]) + [
                                {**contests(Republican=[])[0],
                                 "office": "STATE SENATOR DISTRICT 3"}],
                            2026)


def test_ballot_gate_rejects_two_candidates_of_one_party(make_db, tmp_path):
    conn = make_db(tmp_path / "wi.sqlite")
    conn.execute("INSERT INTO people (id, name, party) VALUES ('p', 'Pat Seat', 'Democratic')")
    conn.execute("INSERT INTO meta VALUES ('wec_ballot_phase', 'general')")
    conn.execute("INSERT INTO elections (person_id, cycle_year, on_ballot, opponents_json, source)"
                 " VALUES ('p', 2026, 1, ?, 'wec')",
                 (json.dumps([{"name": "Ann Lee", "party": "Democratic"}]),))
    conn.execute("INSERT INTO statewide_races (office, incumbent_noncandidacy, candidate, party,"
                 " source) VALUES ('GOVERNOR', 0, 'A', 'Republican', 'wec'),"
                 " ('GOVERNOR', 0, 'B', 'Republican', 'wec')")
    assert len(check_ballot(conn)) == 2
    conn.execute("UPDATE meta SET value = 'primary'")
    assert check_ballot(conn) == [], "before the primary several filings per party are normal"
    conn.close()
